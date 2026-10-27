from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from openpyxl import load_workbook

from kmlog.classify import Answer, apply_answer
from kmlog.cli import main
from kmlog.model import (
    Classification,
    DayNote,
    Leg,
    LegSource,
    LegStatus,
    Rate,
    Vehicle,
    VehicleKind,
)
from kmlog.report.maandstaat import (
    build_jaaroverzicht,
    build_maandstaat,
    write_jaaroverzicht,
    write_maandstaat,
)
from kmlog.store.sqlite import SqliteStore
from test_stamdata import EXAMPLE
from test_voorbeelddagen import (
    LOCATIONS,
    WORKDAY,
    _day_c,
    _private_day,
    normal_day,
    observations,
    run,
)

CAR = Vehicle("car", "BMW i4", "AB-123-C", VehicleKind.CAR, True)
RATE = Rate(2026, 25)


def manual(day: date, km: int, note: str | None = None) -> Leg:
    return Leg(
        id=f"car#manual#{day}",
        vehicle="car",
        date=day,
        distance_km=Decimal(km),
        business_km=Decimal(km),
        classification=Classification.BUSINESS,
        status=LegStatus.CONFIRMED,
        note=note,
        source=LegSource.MANUAL,
    )


def test_maandstaat_mixes_manual_cardata_and_notes():
    a_legs, _ = run(WORKDAY, normal_day(70))  # 7 okt: 141 km
    c_legs, _ = _day_c()  # also 7 okt in the fixture; move it to 8 okt
    for leg in c_legs:
        leg.date = date(2026, 10, 8)
    obs = observations(WORKDAY, normal_day(70))
    m = build_maandstaat(
        2026,
        10,
        CAR,
        RATE,
        [manual(date(2026, 10, 5), 140), *a_legs, *c_legs, manual(date(2026, 9, 30), 140)],
        [DayNote("car", date(2026, 10, 6), "Motor")],
        LOCATIONS,
        obs,
    )
    assert [(r.date.day, r.description) for r in m.rows] == [
        (5, "handmatige rit"),
        (6, "Motor"),
        (7, "Thuis → Klant → Thuis"),
        (8, "Thuis → Klant → onbekend → Thuis"),
    ]
    day7, day8 = m.rows[2], m.rows[3]
    assert (day7.driven_km, day7.business_km) == (141, 141)
    assert (day7.odometer_from, day7.odometer_to) == (40_000, 40_141)
    assert (m.rows[0].odometer_from, m.rows[0].odometer_to) == (None, None)  # manual trip
    assert (day8.driven_km, day8.business_km, day8.private_km) == (160, 141, 19)
    assert not day8.open_confirmation  # rule 3: stop at an unknown place, no question
    assert m.business_km == 140 + 141 + 141
    assert m.amount_cents == 422 * 25
    assert (m.odometer_start, m.odometer_end) == (40_000, 40_141)
    assert m.open_confirmations == []


def test_answer_changes_maandstaat():
    legs, segments = run(WORKDAY, normal_day(96))
    apply_answer(segments[-1], Answer.DETOUR)
    m = build_maandstaat(2026, 10, CAR, RATE, legs, locations=LOCATIONS)
    assert m.business_km == 167
    assert m.open_confirmations == []


def test_motor_is_never_reimbursed():
    motor = Vehicle("bike", "Motor", "XX-00-X", VehicleKind.MOTOR, reimbursable=False)
    m = build_maandstaat(2026, 10, motor, RATE, [manual(date(2026, 10, 5), 140)])
    assert m.business_km == 140
    assert m.amount_cents == 0


def test_rate_year_must_match():
    with pytest.raises(ValueError):
        build_maandstaat(2027, 1, CAR, RATE, [])


def test_jaaroverzicht_and_xlsx(tmp_path):
    sep = build_maandstaat(2026, 9, CAR, RATE, [manual(date(2026, 9, 14), 88, "Leverancier")])
    okt_legs, _ = run(WORKDAY, normal_day(70))
    okt = build_maandstaat(
        2026, 10, CAR, RATE, [manual(date(2026, 10, 5), 140), *okt_legs], locations=LOCATIONS
    )
    totals = build_jaaroverzicht([okt, sep])
    assert [(t.month, t.trip_days, t.business_km, t.amount_cents) for t in totals] == [
        (9, 1, 88, 2200),
        (10, 2, 281, 7025),
    ]

    path = write_maandstaat(okt, tmp_path / "okt.xlsx")
    ws = openpyxl.load_workbook(path).active
    values = [v for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    assert "Maandstaat oktober 2026" in values
    assert "Thuis → Klant → Thuis" in values
    assert 281 in values

    trips = openpyxl.load_workbook(path)["Ritten oktober 2026"]
    rows = [r for r in trips.iter_rows(min_row=5, values_only=True) if r[0] is not None]
    assert [(r[2], r[3], r[4], r[8], r[9]) for r in rows] == [
        (None, None, "handmatige rit", 140, 140),
        ("07:00", "08:30", "Thuis → Klant", 71, 71),
        ("16:00", "17:45", "Klant → Thuis", 70, 70),
    ]

    path = write_jaaroverzicht(2026, CAR, [sep, okt], tmp_path / "jaar.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Jaaroverzicht 2026", "September 2026", "Oktober 2026"]


def test_cli_stamdata_and_reports(tmp_path, capsys):
    db = tmp_path / "km.sqlite"
    main(["--db", str(db), "stamdata", str(EXAMPLE)])  # vehicle + rate 2026
    store = SqliteStore(db)
    store.put_leg(manual(date(2026, 9, 14), 88))
    store.put_leg(manual(date(2026, 9, 16), 140))
    store.close()
    main(["--db", str(db), "maandstaat", "2026", "9", "--out", str(tmp_path / "sep.xlsx")])
    main(["--db", str(db), "jaaroverzicht", "2026", "--out", str(tmp_path / "jaar.xlsx")])
    out = capsys.readouterr().out
    assert "Maandstaat september 2026: 228 km = € 57,00" in out
    assert "Totaal 2026: 228 km = € 57,00" in out
    assert (tmp_path / "sep.xlsx").exists() and (tmp_path / "jaar.xlsx").exists()


def test_trips_one_row_per_leg():
    legs, _ = _day_c()  # client -> unknown -> home: rule 3
    m = build_maandstaat(2026, 10, CAR, RATE, legs, locations=LOCATIONS)
    assert len(m.rows) == 1 and len(m.trips) == 3
    assert [t.description for t in m.trips] == [
        "Thuis → Klant",
        "Klant → onbekend",
        "onbekend → Thuis",
    ]
    assert [t.explanation for t in m.trips] == [
        "Zakelijk: past bij de verwachte 70 km (t/m 80 km).",
        *["Tussenstop: zakelijk telt de verwachte 70 km, de rest is privé."] * 2,
    ]
    assert sum(t.business_km for t in m.trips) == m.business_km


def test_open_question_marks_the_trip():
    legs, _ = run(WORKDAY, normal_day(96))  # rule 2: longer than expected
    m = build_maandstaat(2026, 10, CAR, RATE, legs, locations=LOCATIONS)
    assert [t.open_confirmation for t in m.trips] == [False, True]
    assert m.trips[1].remark == "te bevestigen"
    assert m.open_confirmations == m.rows


def test_review_lists():
    private, _ = _private_day(WORKDAY)  # home -> unknown 23 -> home 23
    short, _ = run(  # 8 km there and back: below the 10 km of the list
        date(2026, 10, 8),
        [
            ("00:00", "10:00", "home", 41_000),
            ("10:40", "14:00", "unknown", 41_004),
            ("14:40", "23:00", "home", 41_008),
        ],
    )
    weekend, _ = run(date(2026, 10, 10), normal_day(70, out_km=70))  # Saturday at the client
    m = build_maandstaat(2026, 10, CAR, RATE, [*private, *short, *weekend], locations=LOCATIONS)
    (place,) = m.unknown_places
    assert [(v.date, v.km) for v in place.visits] == [(WORKDAY, 46)]
    assert place.visits[0].departed_after is not None
    assert [(t.name, t.dates) for t in m.weekend_trips] == [("Klant", [date(2026, 10, 10)])]


def test_cli_version(capsys):
    from kmlog.version import VERSION

    with pytest.raises(SystemExit):
        main(["--version"])
    assert capsys.readouterr().out.strip() == f"kmlog {VERSION}"


def test_private_only_days_and_trips_are_left_out_of_the_excel(tmp_path):
    private, _ = _private_day(WORKDAY)  # home -> unknown -> home, 0 km business
    errand = [  # 3 km home -> home, then a normal client day
        ("00:00", "06:00", "home", 41_000),
        ("06:40", "07:30", "home", 41_003),
        ("08:30", "16:30", "client", 41_074),
        ("17:45", "23:00", "home", 41_144),
    ]
    mixed, _ = run(date(2026, 10, 8), errand)
    first, _ = run(date(2026, 10, 9), normal_day(31, "client2", out_km=75))  # open question
    notes = [DayNote("car", date(2026, 10, 6), "Motor")]
    m = build_maandstaat(2026, 10, CAR, RATE, [*private, *mixed, *first], notes, LOCATIONS)
    # the month page shows everything
    assert [(r.date.day, r.declared) for r in m.rows] == [
        (6, True),
        (7, False),
        (8, True),
        (9, True),
    ]
    assert (m.rows[2].driven_km, m.rows[2].private_km) == (144, 3)  # the day is complete
    assert len(m.trips) == 7

    wb = load_workbook(write_maandstaat(m, tmp_path / "okt.xlsx"))
    table = wb.worksheets[0].iter_rows(min_row=9, max_row=8 + 4, values_only=True)
    days = [r[2] for r in table if isinstance(r[1], str)]  # rows with a weekday
    assert days == ["Motor", "Thuis → Thuis → Klant → Thuis", "Thuis → Klant 2 → Thuis"]
    trips = [r[4] for r in wb.worksheets[1].iter_rows(min_row=5, values_only=True) if r[0]]
    assert trips == ["Thuis → Klant", "Klant → Thuis", "Thuis → Klant 2", "Klant 2 → Thuis"]


def test_clock_shows_the_day_when_it_differs():
    from datetime import UTC, datetime

    from kmlog.report.maandstaat import clock

    ts = datetime(2026, 10, 27, 7, 0, tzinfo=UTC)  # di 27-10 08:00 Amsterdam
    assert clock(ts, date(2026, 10, 27)) == "08:00"
    assert clock(ts, date(2026, 10, 26)) == "di 27-10 08:00"
    assert clock(None, date(2026, 10, 26)) is None
