"""Maandstaat and jaaroverzicht (Excel, later PDF). See ONTWERP.md §9.

Per month: vehicle, rate; per day: date, description, driven / business / private km,
remark (note, flags, open confirmation); totals (business km, amount in cents), odometer at
begin/end of the month (from observations), and the list of open
confirmations. The day rows are the declaration; `trips` lists the same legs one per row
(rit) for inspection, with times, odometer and pass-through points. Open confirmations
already carry their conservative default, so they are counted as-is. Days and trips that are
only private (no business km, no open question) are marked `declared=False`: the web UI shows
them, the Excel (the declaration) leaves them out. `unknown_places` and
`weekend_trips` are the review lists of the monthly mail (kmlog.review). The jaaroverzicht shows
km and amount per month.

Building (pure, testable) and writing (openpyxl) are separate steps.
"""

from __future__ import annotations

import calendar
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..classify import explain
from ..derive import AMSTERDAM, amsterdam_date
from ..model import DayNote, Leg, LegSource, LegStatus, Location, Observation, Rate, Vehicle
from ..review import UnknownPlace, WeekendTrips, unknown_places, weekend_client_trips

MONTHS = [
    "",
    "januari",
    "februari",
    "maart",
    "april",
    "mei",
    "juni",
    "juli",
    "augustus",
    "september",
    "oktober",
    "november",
    "december",
]
WEEKDAYS = ["ma", "di", "wo", "do", "vr", "za", "zo"]
OPEN = (LegStatus.NEEDS_CONFIRMATION,)
FLAG_TEXT = {
    "weekend_or_holiday": "weekend/feestdag",
    "no_route": "geen bekende route",  # before v0.3.0
    "implausible": "eerste rit niet aannemelijk",
    "manual_correction": "handmatig gecorrigeerd",
}


@dataclass
class DayRow:
    date: date
    description: str
    driven_km: int = 0
    business_km: int = 0
    private_km: int = 0
    remark: str = ""
    open_confirmation: bool = False
    declared: bool = True  # False: only private trips, left out of the Excel
    odometer_from: int | None = None  # km-stand when leaving for the first leg of the day
    odometer_to: int | None = None  # km-stand after the last leg of the day


@dataclass
class TripRow:
    """One leg (rit). departed_after is the last poll that saw the car parked, so the real
    departure lies between it and arrived_at; manual legs have no times."""

    date: date
    description: str
    leg_id: str
    departed_after: datetime | None = None
    arrived_at: datetime | None = None
    explanation: str = ""
    driven_km: int = 0
    business_km: int = 0
    private_km: int = 0
    remark: str = ""
    open_confirmation: bool = False
    declared: bool = True
    odometer_from: int | None = None
    odometer_to: int | None = None


@dataclass
class Maandstaat:
    vehicle: Vehicle
    year: int
    month: int
    cents_per_km: int
    rows: list[DayRow] = field(default_factory=list)
    trips: list[TripRow] = field(default_factory=list)
    odometer_start: int | None = None
    odometer_end: int | None = None
    unknown_places: list[UnknownPlace] = field(default_factory=list)
    weekend_trips: list[WeekendTrips] = field(default_factory=list)

    @property
    def title(self) -> str:
        return f"Maandstaat {MONTHS[self.month]} {self.year}"

    @property
    def business_km(self) -> int:
        return sum(r.business_km for r in self.rows)

    @property
    def amount_cents(self) -> int:
        if not self.vehicle.reimbursable:
            return 0
        return Rate(self.year, self.cents_per_km).amount_cents(self.business_km)

    @property
    def open_confirmations(self) -> list[DayRow]:
        return [r for r in self.rows if r.open_confirmation]


def _km(value: Decimal) -> int:
    if value != value.to_integral_value():
        raise ValueError(f"non-integer km {value}")
    return int(value)


def _place(loc_id: str | None, names: dict[str, str]) -> str:
    return names.get(loc_id or "", "onbekend")


def describe(legs: list[Leg], names: dict[str, str]) -> str:
    if all(leg.source is LegSource.MANUAL for leg in legs):
        notes = [leg.note for leg in legs if leg.note]
        if notes:
            return "; ".join(notes)
        return "geïmporteerde rit" if all(leg.imported for leg in legs) else "handmatige rit"
    parts = [_place(legs[0].from_location, names)]
    for leg in legs:
        parts.append(_place(leg.to_location, names))
    return " → ".join(parts)


def _remarks(legs: list[Leg], day_note: str | None = None) -> tuple[str, bool]:
    """Remark (open confirmation, flags, day note, leg notes) and whether a question is open."""
    open_ = any(leg.status in OPEN for leg in legs)
    parts = ["te bevestigen"] if open_ else []
    parts += [FLAG_TEXT.get(f, f) for f in sorted({f for leg in legs for f in leg.flags})]
    if day_note:
        parts.append(day_note)
    if not all(leg.source is LegSource.MANUAL for leg in legs):  # else the note is the description
        parts += [leg.note for leg in legs if leg.note]
    return "; ".join(parts), open_


def _declared(legs: list[Leg]) -> bool:
    """Belongs on the statement: business km, or a question that is still open."""
    return any(leg.business_km or leg.status in OPEN for leg in legs)


def _km_label(km: int) -> str:
    return f"{km:,}".replace(",", ".")


def build_maandstaat(
    year: int,
    month: int,
    vehicle: Vehicle,
    rate: Rate,
    legs: Iterable[Leg],
    day_notes: Iterable[DayNote] = (),
    locations: Iterable[Location] = (),
    observations: Iterable[Observation] = (),
) -> Maandstaat:
    if rate.year != year:
        raise ValueError(f"rate is for {rate.year}, not {year}")
    locations = list(locations)
    names = {loc.id: loc.name for loc in locations}
    first, last = date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])

    by_day: dict[date, list[Leg]] = {}

    def order(leg: Leg) -> tuple:
        departed = leg.departed_after.timestamp() if leg.departed_after else 0.0
        return leg.date, departed, leg.id

    for leg in sorted(legs, key=order):
        if first <= leg.date <= last:
            by_day.setdefault(leg.date, []).append(leg)
    notes = {n.date: n.text for n in day_notes if first <= n.date <= last}

    rows: list[DayRow] = []
    for day in sorted(set(by_day) | set(notes)):
        day_legs = by_day.get(day, [])
        if not day_legs:
            rows.append(DayRow(date=day, description=notes[day]))
            continue
        remark, open_ = _remarks(day_legs, notes.get(day))
        rows.append(
            DayRow(
                date=day,
                description=describe(day_legs, names),
                driven_km=sum(_km(leg.distance_km) for leg in day_legs),
                business_km=sum(_km(leg.business_km) for leg in day_legs),
                private_km=sum(_km(leg.private_km) for leg in day_legs),
                remark=remark,
                open_confirmation=open_,
                declared=_declared(day_legs),
                odometer_from=next(
                    (leg.odometer_from for leg in day_legs if leg.odometer_from is not None), None
                ),
                odometer_to=next(
                    (leg.odometer_to for leg in reversed(day_legs) if leg.odometer_to is not None),
                    None,
                ),
            )
        )

    trips: list[TripRow] = []
    for day in sorted(by_day):
        for leg in by_day[day]:
            remark, open_ = _remarks([leg])
            trips.append(
                TripRow(
                    date=day,
                    description=describe([leg], names),
                    leg_id=leg.id,
                    departed_after=leg.departed_after,
                    arrived_at=leg.arrived_at,
                    explanation=explain(leg),
                    driven_km=_km(leg.distance_km),
                    business_km=_km(leg.business_km),
                    private_km=_km(leg.private_km),
                    remark=remark,
                    open_confirmation=open_,
                    declared=_declared([leg]),
                    odometer_from=leg.odometer_from,
                    odometer_to=leg.odometer_to,
                )
            )

    odometers = [
        o.odometer_km for o in observations if first <= amsterdam_date(o.polled_at) <= last
    ]
    month_legs = [leg for day_legs in by_day.values() for leg in day_legs]
    return Maandstaat(
        vehicle=vehicle,
        year=year,
        month=month,
        cents_per_km=rate.cents_per_km,
        rows=rows,
        trips=trips,
        odometer_start=min(odometers, default=None),
        odometer_end=max(odometers, default=None),
        unknown_places=unknown_places(month_legs),
        weekend_trips=weekend_client_trips(month_legs, locations),
    )


@dataclass
class MonthTotal:
    month: int
    trip_days: int
    business_km: int
    amount_cents: int


def build_jaaroverzicht(maandstaten: Iterable[Maandstaat]) -> list[MonthTotal]:
    return [
        MonthTotal(
            month=m.month,
            trip_days=sum(1 for r in m.rows if r.business_km),
            business_km=m.business_km,
            amount_cents=m.amount_cents,
        )
        for m in sorted(maandstaten, key=lambda m: m.month)
    ]


# --- writing ------------------------------------------------------------------------------

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HEADER_FILL = PatternFill("solid", fgColor="DDE5F0")
OPEN_FILL = PatternFill("solid", fgColor="FFF2CC")
EURO = '"€" #,##0.00'


def format_euro(cents: int) -> str:
    """Dutch notation: € 2.452,75"""
    sign = "-" if cents < 0 else ""
    euros, rest = divmod(abs(cents), 100)
    return f"{sign}€ {euros:,}".replace(",", ".") + f",{rest:02d}"


def _euros(cents: int) -> Decimal:
    return Decimal(cents) / 100


def _table_header(ws, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row, col, text)
        cell.font, cell.fill = BOLD, HEADER_FILL


def _widths(ws, widths: list[int]) -> None:
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def write_maandstaat_sheet(ws, m: Maandstaat) -> None:
    ws.title = f"{MONTHS[m.month].capitalize()} {m.year}"
    ws["A1"] = m.title
    ws["A1"].font = TITLE
    meta = [
        ("Voertuig", f"{m.vehicle.name} ({m.vehicle.plate})"),
        ("Tarief", _euros(m.cents_per_km)),
        ("Kilometerstand begin", m.odometer_start if m.odometer_start is not None else "–"),
        ("Kilometerstand einde", m.odometer_end if m.odometer_end is not None else "–"),
    ]
    for i, (label, value) in enumerate(meta, 3):
        ws.cell(i, 1, label).font = BOLD
        cell = ws.cell(i, 2, value)
        if label == "Tarief":
            cell.number_format = EURO
        cell.alignment = Alignment(horizontal="left")

    head = 8
    _table_header(
        ws,
        head,
        [
            "Datum",
            "Dag",
            "Omschrijving",
            "Km-stand vertrek",
            "Km-stand aankomst",
            "Gereden km",
            "Zakelijk km",
            "Privé km",
            "Opmerking",
        ],
    )
    row = head
    for r in (r for r in m.rows if r.declared):
        row += 1
        values = [r.date, WEEKDAYS[r.date.weekday()], r.description]
        values += [r.odometer_from, r.odometer_to]
        if r.business_km or r.driven_km:
            values += [r.driven_km, r.business_km, r.private_km]
        else:
            values += [None, None, None]
        values.append(r.remark)
        for col, v in enumerate(values, 1):
            cell = ws.cell(row, col, v)
            if r.open_confirmation:
                cell.fill = OPEN_FILL
        ws.cell(row, 1).number_format = "DD-MM-YYYY"

    row += 2
    ws.cell(row, 3, "Totaal zakelijk km").font = BOLD
    ws.cell(row, 7, m.business_km).font = BOLD
    row += 1
    ws.cell(row, 3, "Te declareren").font = BOLD
    amount = ws.cell(row, 7, _euros(m.amount_cents))
    amount.font, amount.number_format = BOLD, EURO
    if not m.vehicle.reimbursable:
        ws.cell(row, 9, "voertuig van de BV: geen km-vergoeding")

    if m.open_confirmations:
        row += 2
        ws.cell(row, 1, "Openstaande bevestigingen (tellen mee met de conservatieve default)")
        ws.cell(row, 1).font = BOLD
        for r in m.open_confirmations:
            row += 1
            ws.cell(row, 1, r.date).number_format = "DD-MM-YYYY"
            ws.cell(row, 3, r.description)
            ws.cell(row, 7, r.business_km)
            ws.cell(row, 9, r.remark)
    _widths(ws, [12, 5, 40, 15, 17, 11, 11, 9, 40])
    ws.freeze_panes = ws.cell(head + 1, 1)


def _hhmm(ts: datetime | None) -> str | None:
    return ts.astimezone(AMSTERDAM).strftime("%H:%M") if ts else None


def write_trips_sheet(ws, m: Maandstaat) -> None:
    """Every declared leg of the month on its own row, for inspection; the declaration is the
    day sheet. Private-only legs are left out, as on the day sheet."""
    ws.title = f"Ritten {MONTHS[m.month]} {m.year}"
    ws["A1"] = f"Ritten {MONTHS[m.month]} {m.year}"
    ws["A1"].font = TITLE
    ws["A2"] = (
        "Vertrokken na = laatste meting waarin de auto nog stond; het echte vertrek ligt "
        "tussen dat moment en de aankomst. Telt niet anders dan de maandstaat."
    )
    head = 4
    _table_header(
        ws,
        head,
        [
            "Datum",
            "Dag",
            "Vertrokken na",
            "Aangekomen",
            "Rit",
            "Uitleg",
            "Km-stand vertrek",
            "Km-stand aankomst",
            "Gereden km",
            "Zakelijk km",
            "Privé km",
            "Opmerking",
        ],
    )
    trips = [t for t in m.trips if t.declared]
    row = head
    for t in trips:
        row += 1
        values = [
            t.date,
            WEEKDAYS[t.date.weekday()],
            _hhmm(t.departed_after),
            _hhmm(t.arrived_at),
            t.description,
            t.explanation,
            t.odometer_from,
            t.odometer_to,
            t.driven_km,
            t.business_km,
            t.private_km,
            t.remark,
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row, col, v)
            if t.open_confirmation:
                cell.fill = OPEN_FILL
        ws.cell(row, 1).number_format = "DD-MM-YYYY"
    row += 2
    ws.cell(row, 5, "Totaal").font = BOLD
    for col, value in [
        (9, sum(t.driven_km for t in trips)),
        (10, sum(t.business_km for t in trips)),
        (11, sum(t.private_km for t in trips)),
    ]:
        ws.cell(row, col, value).font = BOLD
    _widths(ws, [12, 5, 13, 12, 30, 50, 15, 17, 11, 11, 9, 40])
    ws.freeze_panes = ws.cell(head + 1, 1)


def write_jaaroverzicht_sheet(ws, year: int, vehicle: Vehicle, totals: list[MonthTotal]) -> None:
    ws.title = f"Jaaroverzicht {year}"
    ws["A1"] = f"Jaaroverzicht {year} — {vehicle.name} ({vehicle.plate})"
    ws["A1"].font = TITLE
    head = 3
    _table_header(ws, head, ["Maand", "Ritdagen", "Zakelijk km", "Bedrag"])
    row = head
    for t in totals:
        row += 1
        ws.cell(row, 1, MONTHS[t.month])
        ws.cell(row, 2, t.trip_days)
        ws.cell(row, 3, t.business_km)
        ws.cell(row, 4, _euros(t.amount_cents)).number_format = EURO
    row += 1
    ws.cell(row, 1, "Totaal").font = BOLD
    for col, value in [
        (2, sum(t.trip_days for t in totals)),
        (3, sum(t.business_km for t in totals)),
    ]:
        ws.cell(row, col, value).font = BOLD
    total = ws.cell(row, 4, _euros(sum(t.amount_cents for t in totals)))
    total.font, total.number_format = BOLD, EURO
    _widths(ws, [14, 10, 12, 12])


def write_maandstaat(m: Maandstaat, out_path: str | Path) -> Path:
    """The maandstaat (declaration) and, if there are any, its trips on a second sheet."""
    wb = Workbook()
    write_maandstaat_sheet(wb.active, m)
    if any(t.declared for t in m.trips):
        write_trips_sheet(wb.create_sheet(), m)
    return _save(wb, out_path)


def write_jaaroverzicht(
    year: int, vehicle: Vehicle, maandstaten: list[Maandstaat], out_path: str | Path
) -> Path:
    """Overview sheet first, followed by one sheet per month."""
    wb = Workbook()
    write_jaaroverzicht_sheet(wb.active, year, vehicle, build_jaaroverzicht(maandstaten))
    for m in sorted(maandstaten, key=lambda m: m.month):
        write_maandstaat_sheet(wb.create_sheet(), m)
    return _save(wb, out_path)


def _save(wb: Workbook, out_path: str | Path) -> Path:
    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
